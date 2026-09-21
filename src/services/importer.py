"""Leitura tolerante de planilhas e CSVs exportados por ferramentas comuns."""

import csv
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from database import DB_PATH, adicionar_tarefa, atualizar_status


@dataclass(frozen=True)
class ImportResult:
    importadas: int
    ignoradas: int
    avisos: tuple[str, ...]


ALIASES = {
    "titulo": ("titulo", "tarefa", "task", "nome", "title"),
    "descricao": ("descricao", "detalhes", "description", "notes", "observacoes"),
    "categoria": ("categoria", "category", "grupo"),
    "prioridade": ("prioridade", "priority"),
    "status": ("status", "situacao", "estado"),
    "data_limite": ("data_limite", "prazo", "data", "due_date", "deadline", "date"),
    "hora_limite": ("hora_limite", "hora", "time", "due_time"),
}


def importar_planilha(caminho: str | Path, db_path: Path | str = DB_PATH) -> ImportResult:
    path = Path(caminho)
    if path.suffix.lower() == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as arquivo:
            return _importar_registros(csv.DictReader(arquivo), db_path)
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("Selecione um arquivo .xlsx, .xlsm ou .csv.")
    try:
        from openpyxl import load_workbook
    except ImportError as erro:
        raise ImportError("OpenPyXL não está instalado.") from erro
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    try:
        cabecalhos = next(rows)
    except StopIteration:
        return ImportResult(0, 0, ("A planilha está vazia.",))
    chaves = [_normalizar_cabecalho(valor) for valor in cabecalhos]
    registros = (dict(zip(chaves, valores, strict=False)) for valores in rows)
    return _importar_registros(registros, db_path)


def _importar_registros(registros: Iterable[dict[str, Any]], db_path: Path | str = DB_PATH) -> ImportResult:
    importadas = 0
    ignoradas = 0
    avisos: list[str] = []
    for indice, registro in enumerate(registros, start=2):
        normalizado = {_normalizar_cabecalho(chave): valor for chave, valor in registro.items() if chave is not None}
        titulo = _valor(normalizado, "titulo")
        if not titulo:
            ignoradas += 1
            if len(avisos) < 5:
                avisos.append(f"Linha {indice}: título ausente.")
            continue
        data_limite = _formatar_data(_valor(normalizado, "data_limite"))
        hora_limite = _formatar_hora(_valor(normalizado, "hora_limite"))
        prioridade = str(_valor(normalizado, "prioridade") or "Media").capitalize()
        if prioridade not in {"Baixa", "Media", "Alta"}:
            prioridade = "Media"
        tarefa_id = adicionar_tarefa(
            str(titulo),
            str(_valor(normalizado, "descricao") or ""),
            str(_valor(normalizado, "categoria") or "Geral"),
            prioridade,
            data_limite,
            db_path=db_path,
            hora_limite=hora_limite,
        )
        status = str(_valor(normalizado, "status") or "").casefold()
        if status in {"concluida", "concluído", "concluída", "completed", "done"}:
            atualizar_status(tarefa_id, "Concluida", db_path)
        importadas += 1
    return ImportResult(importadas, ignoradas, tuple(avisos))


def _normalizar_cabecalho(valor: object) -> str:
    texto = "" if valor is None else str(valor).strip().casefold().replace(" ", "_")
    return "".join(caractere for caractere in unicodedata.normalize("NFD", texto) if unicodedata.category(caractere) != "Mn")


def _valor(registro: dict[str, Any], campo: str) -> Any:
    for alias in ALIASES[campo]:
        if alias in registro and registro[alias] not in (None, ""):
            return registro[alias]
    return None


def _formatar_data(valor: object) -> str | None:
    if not valor:
        return None
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    texto = str(valor).strip()
    for formato in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(texto, formato).date().isoformat()
        except ValueError:
            continue
    return None


def _formatar_hora(valor: object) -> str | None:
    if not valor:
        return None
    if isinstance(valor, datetime):
        return valor.strftime("%H:%M")
    texto = str(valor).strip()
    for formato in ("%H:%M", "%H:%M:%S"):
        try:
            return datetime.strptime(texto, formato).strftime("%H:%M")
        except ValueError:
            continue
    return None
